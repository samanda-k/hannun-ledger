"""
월별 수입지출 엑셀 → 한눈 가계부(구글 시트) 가져오기

사용법 (PowerShell):
  python tools/import_excel.py "D:/AI자동화/가계부_가져오기/월별수입지출내역_2025.xlsx"          # 확인만 (시트에 안 넣음)
  python tools/import_excel.py "...xlsx" --push                                                  # 시트에 넣기

엑셀은 해마다 모양이 조금씩 달라서, 칸 위치는 머리줄 글자(결제구분·적요·공급가·부가세·입금액…)로 찾아요.
  - 매출: 첫 '날짜' 머리줄 아래. 날짜·인원은 그날 첫 줄에만. 결제구분이 '대출'이면 대출 받은 돈
  - '모으기' / '기타수입' 구역
  - '지출' → '고정' / '사업자' / '개인' 구역 (2023년처럼 구역 없이 한 목록인 해도 있음)
  - 비고(구매처) 칸의 글자로 분류: 고정비·소모품·이자비용·대출원금·보험료·초기비용·이전·보증금·슈가·제품·알토란·광고비…
연결 정보는 엑셀과 같은 폴더의 연결정보.txt (URL=..., CODE=...) 에서 읽어요. 화면에 출력하지 않아요.
"""
import argparse
import calendar
import datetime as dt
import hashlib
import json
import os
import re
import sys
import time
import warnings
from collections import defaultdict

import openpyxl
import requests

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')

SECTION_MARKERS = {'고정': 'fix', '사업자': 'biz', '개인': 'per'}
AMOUNT_HEADS = ('입금액', '판매금액', '지출액', '금액')
NOTE_HEADS = ('비고', '구매처')
PERSONAL_TAGS = {'식대': '식비', '여행': '여행', '운동': '운동'}
SETUP_NOTES = ('초기비용', '이전', '보증금', '인테리어')   # 창업·이전 비용 (월 급여·남는 돈 계산에서 빠짐)
FUND_CAT = '대출·자금'                                  # 대출·이전 자금·창업 부가세 환급 (계산에서 빠짐)
BIZ_ETC_NOTES = ('알토란', '슈가', '제품', '광고비', '면허세', '세금', '차량유지비')


def s(v):
    return '' if v is None else str(v).strip()


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', ''))
    except (TypeError, ValueError):
        return None


def is_date(v):
    return isinstance(v, (dt.datetime, dt.date))


def header_cols(row, below=()):
    """머리줄 → {역할: 열 번호(0부터)}. 결제구분·인원 칸은 머리글이 어긋난 달이 있어서 아래 실제 값으로 정해요"""
    cols = {}
    labels = [s(v) for v in row]
    for i, lab in enumerate(labels):
        if lab == '인원':
            cols['vis'] = i
        elif lab == '적요':
            cols['item'] = i
        elif lab == '공급가':
            cols['sup'] = i
        elif lab == '부가세':
            cols['vat'] = i
        elif lab == '일 합계':
            cols['day'] = i
        elif lab in AMOUNT_HEADS and 'amt' not in cols:
            cols['amt'] = i
        elif lab in NOTE_HEADS:
            cols['note'] = i
    if 'note' not in cols and 'amt' in cols and 'day' not in cols:
        cols['note'] = cols['amt'] + 1        # 구매처 머리글이 빈 달 (2023년 6·7월)
    if 'item' in cols:
        cand = range(1, cols['item'])
        texts = {c: 0 for c in cand}
        nums = {c: 0 for c in cand}
        for r in below[:60]:
            if s(r[0]) in ('모으기', '지출', '고정', '사업자', '개인'):
                break
            for c in cand:
                v = r[c] if c < len(r) else None
                if isinstance(v, str) and v.strip():
                    texts[c] += 1
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    nums[c] += 1
        best = max(cand, key=lambda c: (texts[c], c)) if cand else None
        if best is not None and texts[best] > 0:
            cols['pay'] = best
        else:
            pays = [i for i, lab in enumerate(labels) if lab == '결제구분' and i < cols['item']]
            cols['pay'] = pays[-1] if pays else cols['item'] - 1
        # 인원: 머리글이 없거나 어긋나도, 결제구분 왼쪽 칸에 숫자가 있으면 그 칸
        if cols['pay'] >= 2 and nums.get(1, 0) > 0:
            cols['vis'] = 1
        elif cols.get('vis') == cols['pay']:
            del cols['vis']
    return cols


def classify_expense(sec, item, note, tag):
    """지출 한 줄 → (t, biz, cat, memo)"""
    memo = item
    if '새출발' in note or '새출발' in item:
        return 'move', False, '새출발기금', memo
    if note in SETUP_NOTES:
        return 'out', True, '창업·이전', memo + (f' ({note})' if note != '초기비용' else '')
    if note == '대출원금' or (note == '' and '원금' in item and sec != 'per'):
        return 'move', False, '빌린 돈 갚기', memo
    if note == '이자비용':
        return 'out', False, '이자', memo
    if note == '고정비' or (note == '보험료' and sec != 'per'):
        return 'out', True, '고정비', memo
    if note == '소모품' or tag == '소모품':
        return 'out', True, '소모품', memo
    if note in BIZ_ETC_NOTES:
        return 'out', True, '사업자', memo + f' ({note})'
    if sec == 'fix':
        return 'out', True, '고정비', memo + (f' ({tag})' if tag else '')
    if sec in ('biz', 'flat'):
        return 'out', True, '사업자', memo + (f' ({note})' if note else '')
    return 'out', False, PERSONAL_TAGS.get(tag, '개인지출'), memo


def parse_month(ws, year, month):
    """한 달 시트 → (entries, visits, check)"""
    entries, visits = [], {}
    check = defaultdict(float)
    last = calendar.monthrange(year, month)[1]
    first_day = f'{year}-{month:02d}-01'
    mode, cur_date, sc, ec = None, None, {}, {}
    width = max(ws.max_column, 10)
    rows = [list(r) + [None] * width for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=width, values_only=True)]
    for i, r in enumerate(rows, start=1):
        a = r[0]
        A = s(a)

        # ── 구역 바뀜 ──
        if A == '날짜' and mode is None:
            sc, mode, cur_date = header_cols(r, rows[i:]), 'sales', None
            continue
        if A == '모으기':
            mode, cur_date = 'save', None
            continue
        if A.replace(' ', '') == '기타수입':
            mode, cur_date = 'etcin', None
            continue
        if A == '지출':
            mode, cur_date = 'flat', None
            continue
        if A == '날짜' and mode in ('flat', 'fix', 'biz', 'per'):
            ec = header_cols(r, rows[i:])
            continue
        if A in SECTION_MARKERS and mode in ('flat', 'fix', 'biz', 'per'):
            mode, cur_date = SECTION_MARKERS[A], None
            check['소계:' + A] += num(r[ec.get('amt', 6)]) or 0
            continue
        if mode is None or mode == 'after_sales':
            continue

        if is_date(a):
            # 지난달 시트를 복사하면서 남은 옛 날짜도 이 시트의 달로 맞춰요 (일은 그대로)
            cur_date = f'{year}-{month:02d}-{min(a.day, last):02d}'
        cols = sc if mode in ('sales', 'save', 'etcin') else ec
        get = lambda k: r[cols[k]] if k in cols else None
        amt, pay, item = num(get('amt')), s(get('pay')), s(get('item'))
        date = cur_date or first_day

        if mode == 'sales':
            is_total = amt and not pay and not item and (
                (a is not None and not is_date(a)) or num(get('sup')) is not None or num(get('day')) is not None)
            if (a is not None and not is_date(a) and not isinstance(a, str)) or is_total:
                check['매출합계'] = num(get('day')) or amt or 0
                if 'vis' in sc:
                    check['방문합계'] = num(r[sc['vis']]) or 0
                mode = 'after_sales'
                continue
            if is_date(a) and 'vis' in sc:
                n = num(r[sc['vis']])
                if n:
                    visits[date] = visits.get(date, 0) + int(round(n))
            if not amt:
                continue
            total = int(round(amt))
            if pay == '대출' or '환급' in pay or '환급' in item:
                entries.append(dict(d=date, t='in', biz=False, cat=FUND_CAT, pay='', sup=0, vat=0, amt=total,
                                    memo=(item + (f' ({pay})' if pay != '대출' else '')).strip(), src=f'{ws.title}!{i}'))
                continue
            if not pay and not item:
                pay, item = '계좌', '시술'     # 결제구분·적요가 빈 매출 줄 (사장님 확인: 계좌 시술)
                check['빈매출줄'] += total
            sup = num(get('sup'))
            sup = int(round(sup)) if sup else 0
            entries.append(dict(d=date, t='in', biz=True, cat=item or '기타', pay=pay, sup=sup,
                                vat=(total - sup) if sup else 0, amt=total, memo='', src=f'{ws.title}!{i}'))
            check['매출'] += total
            continue

        if not amt or not (item or (mode in ('fix', 'biz', 'per', 'flat') and pay)):
            continue                  # 합계 줄·예상치 줄(적요도 결제구분도 없음)은 건너뛰어요
        total = int(round(amt))
        if not item:
            item = pay                # 적요만 비고 결제구분은 있는 지출 줄

        if mode == 'save':
            if '빌린' in item:
                cat, memo = '빌린 돈 갚기', pay or item
            elif '노란우산' in item:
                cat, memo = '노란우산공제', item + (f' ({pay})' if pay else '')
            else:
                cat, memo = '모으기', item + (f' ({pay})' if pay else '')
            entries.append(dict(d=date, t='move', biz=False, cat=cat, pay='', sup=0, vat=0, amt=total,
                                memo=memo, src=f'{ws.title}!{i}'))
            continue

        if mode == 'etcin':
            entries.append(dict(d=date, t='in', biz=False, cat='기타 수입', pay='', sup=0, vat=0, amt=total,
                                memo=item, src=f'{ws.title}!{i}'))
            continue

        # 지출
        note = s(get('note'))
        t, biz, cat, memo = classify_expense(mode, item, note, pay)
        sup = num(get('sup'))
        sup = int(round(sup)) if sup and t == 'out' else 0
        entries.append(dict(d=date, t=t, biz=biz, cat=cat, pay='', sup=sup, vat=(total - sup) if sup else 0,
                            amt=total, memo=memo, src=f'{ws.title}!{i}', paid_by=pay))
        check['구역합:' + {'fix': '고정', 'biz': '사업자', 'per': '개인', 'flat': '지출'}[mode]] += total
    return entries, visits, check


def assign_ids(entries):
    seen = defaultdict(int)
    for e in entries:
        key = '|'.join(str(e[k]) for k in ('d', 't', 'biz', 'cat', 'pay', 'amt', 'memo'))
        seen[key] += 1
        e['id'] = 'xl' + hashlib.sha1(f'{key}|{seen[key]}'.encode()).hexdigest()[:12]
        e['ts'] = int(time.mktime(time.strptime(e['d'], '%Y-%m-%d')) * 1000) + seen[key]


KEYS = ['수입', '방문', '고정비', '소모품', '사업자', '지출계', '월급여', '기타수입', '개인지출', '빌린돈', '새출발', '모으기',
        '노란우산', '남는돈', '창업이전', '대출받음']


def month_stats(entries, visits, ym):
    r = {k: 0 for k in KEYS}
    for e in entries:
        if not e['d'].startswith(ym):
            continue
        c, a = e['cat'], e['amt']
        if e['t'] == 'in':
            if e['biz']:
                r['수입'] += a
            elif c == FUND_CAT:
                r['대출받음'] += a
            else:
                r['기타수입'] += a
        elif e['t'] == 'out':
            if c == '창업·이전':
                r['창업이전'] += a
            elif e['biz']:
                r[c if c in ('고정비', '소모품') else '사업자'] += a
            else:
                r['개인지출'] += a
        else:
            r[{'새출발기금': '새출발', '모으기': '모으기', '노란우산공제': '노란우산'}.get(c, '빌린돈')] += a
    r['방문'] = sum(n for d, n in visits.items() if d.startswith(ym))
    r['지출계'] = r['고정비'] + r['소모품'] + r['사업자']
    r['월급여'] = r['수입'] - r['지출계']
    r['남는돈'] = r['월급여'] + r['기타수입'] - r['개인지출'] - r['빌린돈'] - r['새출발'] - r['모으기'] - r['노란우산']
    return r


def total_sheet_values(wb, year):
    """'토탈' 시트의 월별 값 (엑셀이 계산해 둔 값)"""
    ws = wb.worksheets[0]
    labels = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=14, values_only=True):
        lab = s(row[0])
        if lab and lab not in labels:
            labels[lab] = [num(v) or 0 for v in row[1:13]]
    pick = lambda *names: next((labels[n] for n in names if n in labels), None)
    return {'수입': pick(f'{year}년 수입', '수입 금액'), '방문': pick(f'{year}년 방문', '방문 인원'),
            '기타수입': pick('기타 수입')}


def read_conn(folder):
    path = os.path.join(folder, '연결정보.txt')
    conf = {}
    with open(path, encoding='utf-8') as fh:
        for line in fh:
            if '=' in line:
                k, v = line.split('=', 1)
                conf[k.strip()] = v.strip()
    conf['URL'] = re.sub(r'/macros/u/\d+/s/', '/macros/s/', conf.get('URL', ''))
    if not conf.get('URL', '').startswith('https://script.google.com/') or not conf.get('CODE') or conf['CODE'].startswith('<'):
        sys.exit('연결정보.txt 의 URL / CODE 를 확인해 주세요.')
    return conf['URL'], conf['CODE']


def api(url, token, body):
    r = requests.post(url, data=json.dumps({**body, 'token': token}).encode('utf-8'),
                      headers={'Content-Type': 'text/plain;charset=utf-8'}, timeout=180)
    try:
        j = r.json()
    except ValueError:
        sys.exit(f'서버 응답을 읽지 못했어요 (HTTP {r.status_code}). 웹 앱 배포 권한을 확인해 주세요.')
    if not j.get('ok'):
        sys.exit(f"서버 오류: {j.get('error')} {j.get('message', '')}")
    return j


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('--push', action='store_true', help='구글 시트에 실제로 넣기')
    _t = dt.date.today()
    ap.add_argument('--until', default=f'{_t.year}-{_t.month:02d}-{calendar.monthrange(_t.year, _t.month)[1]:02d}',
                    help='이 날짜까지만 가져오기 (기본: 이번 달 말일)')
    ap.add_argument('--months', default='', help='가져올 달 (예: 1-9)')
    ap.add_argument('--skip', action='append', default=[], help='빼고 넣을 것 "월:분류" (예: "7월:빌린 돈 갚기")')
    ap.add_argument('--skip-pay', action='append', default=[], help='결제구분에 이 글자가 있는 지출은 빼기 (예: "(오빠)")')
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    m = re.search(r'(20\d\d)', os.path.basename(args.xlsx))
    year = int(m.group(1)) if m else _t.year
    months = range(1, 13)
    if args.months:
        lo, _, hi = args.months.partition('-')
        months = range(int(lo), int(hi or lo) + 1)

    entries, visits, checks = [], {}, {}
    for mo in months:
        if f'{mo}월' not in wb.sheetnames:
            continue
        e, v, c = parse_month(wb[f'{mo}월'], year, mo)
        entries += e
        visits.update(v)
        checks[mo] = c
    entries = [e for e in entries if e['d'] <= args.until]
    for spec in args.skip:
        sheet, _, cat = spec.partition(':')
        before = len(entries)
        entries = [e for e in entries if not (e['src'].split('!')[0] == sheet and e['cat'] == cat)]
        print(f'빼고 넣기: {sheet} {cat} {before - len(entries)}건')
    for word in args.skip_pay:
        before = len(entries)
        entries = [e for e in entries if not (e['t'] == 'out' and word in e.get('paid_by', ''))]
        print(f'빼고 넣기: 결제구분에 "{word}" 있는 지출 {before - len(entries)}건')
    visits = {d: n for d, n in visits.items() if d <= args.until}
    tot = total_sheet_values(wb, year)

    # 월 시트에 기타수입 구역이 없는 달은 토탈의 '기타 수입' 값을 써요 (대출 받은 달은 제외)
    for mo in months:
        ym = f'{year}-{mo:02d}'
        has = any(e['d'].startswith(ym) and e['t'] == 'in' and not e['biz'] for e in entries)
        v = tot['기타수입'][mo - 1] if tot['기타수입'] else 0
        if v and not has and f'{ym}-01' <= args.until:
            # 창업·이전 비용이 있는 달의 기타 수입은 그 비용을 댄 돈이라 계산에서 빼요 (2025년 5월 이전)
            setup = any(e['d'].startswith(ym) and e['cat'] == '창업·이전' for e in entries)
            entries.append(dict(d=f'{ym}-01', t='in', biz=False, cat=FUND_CAT if setup else '기타 수입', pay='', sup=0, vat=0,
                                amt=int(round(v)), memo='엑셀 토탈 기타 수입' + (' (이전 자금)' if setup else ''), src=f'토탈!{mo}월'))
    assign_ids(entries)

    last_month = max((int(e['d'][5:7]) for e in entries), default=0)
    print(f'\n{year}년 · 기록 {len(entries)}건 · 방문 기록 {len(visits)}일\n')
    show = ['수입', '방문', '지출계', '월급여', '기타수입', '개인지출', '빌린돈', '새출발', '모으기', '노란우산', '남는돈', '창업이전', '대출받음']
    print(f"{'월':>3} " + ' '.join(f'{k:>10}' for k in show))
    notes = []
    sums = defaultdict(int)
    for mo in months:
        if mo > last_month:
            break
        st = month_stats(entries, visits, f'{year}-{mo:02d}')
        for k in show:
            sums[k] += st[k]
        print(f"{mo:>2}월 " + ' '.join(f'{int(st[k]):>10,}' for k in show))
        xs = tot['수입'][mo - 1] if tot['수입'] else None
        if xs is not None and abs(xs - st['수입']) >= 1:
            notes.append(f'{mo}월 매출 {int(st["수입"]):,} ↔ 엑셀 토탈 {int(xs):,}')
        c = checks.get(mo, {})
        if c.get('매출합계') and min(abs(c['매출합계'] - st['수입']), abs(c['매출합계'] - st['수입'] - st['대출받음'])) >= 1:
            notes.append(f'{mo}월 매출 {int(st["수입"]):,} ↔ 월 시트 합계줄 {int(c["매출합계"]):,}')
        if c.get('방문합계') and abs(c['방문합계'] - st['방문']) >= 1:
            notes.append(f'{mo}월 방문 {st["방문"]} ↔ 월 시트 합계줄 {int(c["방문합계"])}')
        for sec in ('고정', '사업자', '개인'):
            sub, got = c.get('소계:' + sec), c.get('구역합:' + sec, 0)
            if sub is not None and abs(sub - got) >= 1:
                notes.append(f'{mo}월 [{sec}] 엑셀 소계 {int(sub):,} ↔ 줄 합계 {int(got):,} (차이 {int(got - sub):+,})')
        if c.get('빈매출줄'):
            notes.append(f'{mo}월 결제구분·적요가 빈 매출 {int(c["빈매출줄"]):,}원 → 계좌 시술로 넣음')
    print(f"{'합계':>2} " + ' '.join(f'{int(sums[k]):>10,}' for k in show))
    cats = defaultdict(int)
    for e in entries:
        if e['t'] == 'out':
            cats[('사업 ' if e['biz'] else '개인 ') + e['cat']] += e['amt']
    print('\n지출 분류 합계:', ', '.join(f'{k} {v:,}' for k, v in sorted(cats.items(), key=lambda x: -x[1])))
    other = [e for e in entries if e['t'] == 'out' and '(' in e.get('paid_by', '')]
    if other:
        notes.append('다른 사람 카드로 낸 지출: ' + ', '.join(f"{e['d'][5:]} {e['paid_by']} {e['memo']} {e['amt']:,}" for e in other))
    print('\n확인할 점:' if notes else '\n확인할 점: 없음')
    for n in notes:
        print('  ' + n)

    if not args.push:
        print('\n(확인만 했어요. 시트에 넣으려면 --push)')
        return

    url, token = read_conn(os.path.dirname(os.path.abspath(args.xlsx)))
    existing = api(url, token, {'action': 'list'})
    have = {e['id'] for e in existing['entries']}
    # 앱에서 이미 입력한 같은 기록(날짜·항목·결제·금액이 같음)은 건너뛰어요
    same = defaultdict(int)
    for e in existing['entries']:
        if not str(e['id']).startswith('xl'):
            same[(e['d'], e['t'], bool(e['biz']), e['cat'], e.get('pay', ''), int(e['amt']))] += 1
    keep, skipped = [], []
    for e in entries:
        k = (e['d'], e['t'], e['biz'], e['cat'], e['pay'], e['amt'])
        if e['id'] not in have and same[k] > 0:
            same[k] -= 1
            skipped.append(e)
        else:
            keep.append(e)
    entries = keep
    if skipped:
        print(f'앱에서 이미 입력한 기록과 같아서 건너뛴 것: {len(skipped)}건')
    ops = [{'op': 'upsert', 'entry': {k: e[k] for k in ('id', 'd', 't', 'biz', 'cat', 'pay', 'sup', 'vat', 'amt', 'memo', 'ts')}}
           for e in entries]
    ops += [{'op': 'visits', 'd': d, 'n': n} for d, n in sorted(visits.items())]
    new = sum(1 for e in entries if e['id'] not in have)
    print(f'\n시트에 넣는 중: 새 기록 {new}건, 이미 있던 기록 {len(entries) - new}건 갱신, 방문 {len(visits)}일')
    CH = 250
    for i in range(0, len(ops), CH):
        api(url, token, {'action': 'batch', 'ops': ops[i:i + CH]})
        print(f'  {min(i + CH, len(ops))}/{len(ops)}')
    after = api(url, token, {'action': 'list'})
    print(f"완료. 시트의 기록 {len(after['entries'])}건 · 방문 {len(after['visits'])}일")


if __name__ == '__main__':
    main()
